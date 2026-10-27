"""Build MM Core as an .xlsx from packages/core/schema.json.

Upload the result to Google Drive with conversion on and you get the Sheet with
every tab, header row, dropdown and seed row. Protected ranges are applied
separately by `npm run sheet:setup` (they need the account emails).

Usage: python3 scripts/build_sheet_xlsx.py [out.xlsx]
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
schema = json.loads((ROOT / "packages/core/schema.json").read_text())
out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "MM Core.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="FFF1E6")   # light orange
LOCKED_FILL = PatternFill("solid", fgColor="EDE9FE")   # light purple = backend-only
HEADER_FONT = Font(bold=True, color="1F1A24")
ROWS = 2000  # validation depth

lists = schema["lists"]
list_names = list(lists.keys())
list_col = {name: get_column_letter(i + 1) for i, name in enumerate(list_names)}

wb = Workbook()
wb.remove(wb.active)

for tab in schema["tabs"]:
    ws = wb.create_sheet(tab["name"])
    if tab["headers"] == "$lists":
        headers = list_names
        for c, name in enumerate(list_names, start=1):
            for r, val in enumerate(lists[name], start=2):
                ws.cell(row=r, column=c, value=val)
    else:
        headers = tab["headers"]

    locked_from = tab.get("backendOnlyFrom")
    locked_idx = headers.index(locked_from) if locked_from else None
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        whole_locked = tab.get("backendOnly") or tab.get("kyleOnly")
        cell.fill = LOCKED_FILL if (whole_locked or (locked_idx is not None and c - 1 >= locked_idx)) else HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(32, len(h) + 4))
    ws.freeze_panes = "A2"

    for r, row in enumerate(tab.get("seed", []), start=2):
        for c, val in enumerate(row, start=1):
            if isinstance(val, bool):
                val = "TRUE" if val else "FALSE"
            ws.cell(row=r, column=c, value=val if val != "" else None)

    for col_name, list_name in tab.get("dropdowns", {}).items():
        c = headers.index(col_name) + 1
        letter = get_column_letter(c)
        n = len(lists[list_name])
        src = f"lists!${list_col[list_name]}$2:${list_col[list_name]}${n + 1}"
        dv = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=True,
                            errorTitle="Not allowed", error=f"Pick a value from lists → {list_name}")
        dv.add(f"{letter}2:{letter}{ROWS}")
        ws.add_data_validation(dv)

wb.save(out)
print(f"wrote {out} with {len(schema['tabs'])} tabs")
